"""Generate the final SR-3 report for LibraryDesk."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from docs.lab08.generate_report import (  # noqa: E402
    add_body,
    add_bullet,
    add_caption,
    add_heading,
    add_image,
    add_listing,
    add_table,
    configure_document,
    new_page,
    set_font,
)


OUTPUT = ROOT / "output" / "docx" / "SR03_Гордєєв_472_2.docx"
EVIDENCE = ROOT / "docs" / "sr03" / "evidence"
LAB8_EVIDENCE = ROOT / "docs" / "lab08" / "evidence"


def create_architecture_diagram():
    """Create a high-resolution dependency diagram for the report."""
    path = EVIDENCE / "architecture.png"
    image = Image.new("RGB", (1800, 950), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    body = ImageFont.truetype(font_path, 34)
    bold = ImageFont.truetype(bold_path, 38)
    small = ImageFont.truetype(font_path, 28)

    boxes = {
        "app": (610, 60, 1190, 210, "LibraryDesk.App", "Console and composition root"),
        "application": (610, 345, 1190, 495, "Application", "Services and ports"),
        "domain": (610, 720, 1190, 870, "Domain", "Entities states and rules"),
        "infra": (70, 345, 500, 495, "Infrastructure", "JSON clock and log"),
        "tests": (1300, 345, 1730, 495, "LibraryDesk.Tests", "Unit and architecture tests"),
    }
    for left, top, right, bottom, title, subtitle in boxes.values():
        draw.rounded_rectangle((left, top, right, bottom), radius=22, fill="#EAF2F8", outline="#17365D", width=5)
        title_width = draw.textbbox((0, 0), title, font=bold)[2]
        subtitle_width = draw.textbbox((0, 0), subtitle, font=small)[2]
        draw.text(((left + right - title_width) / 2, top + 25), title, font=bold, fill="#000000")
        draw.text(((left + right - subtitle_width) / 2, top + 85), subtitle, font=small, fill="#333333")

    def arrow(start, end, label):
        draw.line((start, end), fill="#17365D", width=7)
        x2, y2 = end
        draw.polygon([(x2, y2), (x2 - 18, y2 - 26), (x2 + 18, y2 - 26)], fill="#17365D")
        box = draw.textbbox((0, 0), label, font=body)
        x = (start[0] + end[0] - (box[2] - box[0])) / 2 + 20
        y = (start[1] + end[1]) / 2 - 20
        draw.rectangle((x - 8, y - 5, x + box[2] - box[0] + 8, y + 43), fill="white")
        draw.text((x, y), label, font=body, fill="#17365D")

    arrow((900, 210), (900, 345), "calls")
    arrow((900, 495), (900, 720), "uses")
    arrow((500, 420), (610, 420), "implements ports")
    arrow((1300, 420), (1190, 420), "verifies")
    image.save(path)
    return path


def add_title_page(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(55)
    set_font(paragraph.add_run("ЗВІТ"), "Arial", 24, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run("до самостійної роботи 3"), "Arial", 16, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(38)
    set_font(
        paragraph.add_run("Індивідуальний проєкт конструювання застосунку за варіантом"),
        "Arial",
        14,
        bold=True,
    )

    metadata = [
        "Дисципліна: Конструювання програмного забезпечення",
        "Проєкт: LibraryDesk",
        "Варіант: 2",
        "Виконав: Гордєєв Дмитро Леонідович",
        "Група: 472",
        "Робоча гілка: sr03-individual-project",
        "Версія: 1.0.0",
    ]
    for line in metadata:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)
        set_font(paragraph.add_run(line), "Times New Roman", 12)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(30)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run("https://github.com/kuum-oss/LibraryDesk"), "Courier New", 9.5)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(44)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run("2026"), "Times New Roman", 12)


def add_contents(doc):
    new_page(doc)
    add_heading(doc, "Зміст")
    rows = [
        ("1", "Мета та вихідні дані"),
        ("2", "Специфікація функцій"),
        ("3", "Архітектура рішення"),
        ("4", "Реалізація за практиками курсу"),
        ("5", "Модульні тести та покриття"),
        ("6", "Конвеєр безперервної інтеграції"),
        ("7", "Документація проєкту"),
        ("8", "Журнал рішень і самооцінка"),
        ("9", "Висновки"),
        ("10", "Використані джерела"),
    ]
    add_table(doc, ["Розділ", "Назва"], rows, [1.0, 5.8], 9.5)
    add_heading(doc, "Анотація")
    add_body(
        doc,
        "У роботі навчальний проєкт LibraryDesk доведено до версії 1.0.0. Реалізовано шість сценаріїв, профільне правило строку видачі й пені, JSON-збереження, файловий журнал, консольне меню, архітектурні межі, 97 зелених тестів і локальний еквівалент CI. Рядкове покриття домену та прикладної логіки становить 85,31 %, покриття гілок — 71,62 %.",
    )
    add_body(
        doc,
        "Звіт відокремлює підтверджені локальні результати від дій, що потребують публікації гілки на GitHub. На момент формування документа нова гілка ще не була надіслана, тому зелений віддалений прогон версії 1.0 і pull request мають бути додані після push.",
    )


def add_intro(doc):
    new_page(doc)
    add_heading(doc, "1 Мета та вихідні дані")
    add_body(
        doc,
        "Мета роботи — перетворити накопичений у лабораторних роботах код на завершений інженерний продукт зі специфікацією, свідомою архітектурою, перевіреною бізнес-логікою, автоматичними критеріями якості та документацією для нового користувача.",
    )
    add_heading(doc, "1.1 Вихідний стан", 2)
    add_table(
        doc,
        ["Компонент", "Стан на початок СР-3"],
        [
            ("LibraryDesk.Core", "Сутності Loan, Book, Reader, правила ціни, сховища й звіти з попередніх робіт"),
            ("LibraryDesk.App", "Демонстраційний запуск одного наперед заданого формуляра"),
            ("LibraryDesk.Tests", "54 тести основного рішення"),
            ("Legacy", "12 характеризаційних тестів ізольованого успадкованого модуля"),
            ("CI", "Формат, Release, 66 тестів і поріг покриття 60 %"),
        ],
        [1.7, 5.1],
        8.8,
    )
    add_heading(doc, "1.2 Підсумковий стан", 2)
    add_table(
        doc,
        ["Показник", "Результат"],
        [
            ("Версія", "1.0.0"),
            ("Сценарії", "6 із 6 реалізовано"),
            ("Тести", "97 із 97 пройдено"),
            ("Release", "0 попереджень, 0 помилок"),
            ("Покриття", "85,31 % рядків; 71,62 % гілок"),
            ("Навчальні дані", "50 книг і 20 умовних читачів"),
            ("Локальний quality gate", "./tools/verify.sh завершується успішно"),
        ],
        [2.1, 4.7],
        9.0,
    )


def add_specification(doc):
    new_page(doc)
    add_heading(doc, "2 Специфікація функцій")
    add_body(
        doc,
        "Специфікація в docs/spec.md описує поведінку без прив'язки до меню або класів. Межі системи виключають GUI, вебдоступ, базу даних, авторизацію, платіжний шлюз і роботу з реальними персональними даними.",
    )
    add_table(
        doc,
        ["Входить", "Свідомо не входить"],
        [
            ("Каталог книг і доступні примірники", "Закупівля та бухгалтерський облік фонду"),
            ("Читачі з ознакою активного абонемента", "Ролі працівників і документи читачів"),
            ("Видача, повернення та пеня", "Онлайн-оплата й банківська інтеграція"),
            ("Пошук за читачем і станом", "Зовнішній повнотекстовий каталог"),
            ("Звіт боржників на дату", "GUI, вебдоступ і багатокористувацький режим"),
        ],
        [3.35, 3.45],
        8.5,
    )
    add_heading(doc, "2.1 Сценарії використання", 2)
    add_table(
        doc,
        ["ID", "Сценарій", "Основний результат", "Ключові альтернативи"],
        [
            ("UC-01", "Зареєструвати книгу", "Новий запис або поповнення", "Порожній ISBN; невалідна кількість"),
            ("UC-02", "Зареєструвати читача", "Активний обліковий запис", "Дубль ID; неправильний email"),
            ("UC-03", "Видати книжку", "Active і строк +14 днів", "Немає книги; абонемент неактивний"),
            ("UC-04", "Прийняти повернення", "Returned і сума пені", "Невідомий номер; перехід заборонено"),
            ("UC-05", "Знайти формуляри", "Відсортований список", "Немає критеріїв; немає збігів"),
            ("UC-06", "Боржники на дату", "Рядки звіту", "Порожній звіт; пошкоджене посилання"),
        ],
        [0.7, 1.55, 2.0, 2.55],
        7.8,
    )
    add_body(
        doc,
        "Кожен сценарій має дійову особу, мету, передумову, 4–8 кроків основного перебігу, щонайменше дві альтернативи, постумову та посилання на тести. UC-03 і UC-04 є обов'язковими сценаріями варіанта, UC-06 — обов'язковим звітом.",
    )

    new_page(doc)
    add_heading(doc, "2.2 Профільне правило строку й пені")
    add_body(
        doc,
        "Стандартний строк становить 14 календарних днів. До 7 днів прострочення діє базова денна ставка, з 8-го дня вся сума множиться на 1,5. Пеня округлюється до двох знаків і не перевищує вартості заміни примірника.",
    )
    add_table(
        doc,
        ["Днів", "Денна ставка", "Множник", "Межа", "Результат"],
        [
            ("0", "10,00", "1,0", "1000,00", "0,00"),
            ("1", "10,00", "1,0", "1000,00", "10,00"),
            ("7", "10,00", "1,0", "1000,00", "70,00"),
            ("8", "10,00", "1,5", "1000,00", "120,00"),
            ("10", "10,00", "1,5", "1000,00", "150,00"),
            ("20", "10,00", "1,5", "300,00", "300,00"),
        ],
        [0.8, 1.45, 1.25, 1.45, 1.45],
        8.8,
    )
    add_listing(
        doc,
        "Лістинг 1  Реалізація профільного правила",
        """public decimal Calculate(int overdueDays, decimal dailyFee, decimal bookPrice)
{
    EnsureArgumentsValid(overdueDays, dailyFee, bookPrice);
    decimal rawFee = overdueDays * dailyFee * MultiplierFor(overdueDays);
    return Math.Min(decimal.Round(rawFee, 2), bookPrice);
}""",
        8.4,
    )
    add_heading(doc, "2.3 Життєвий цикл формуляра", 2)
    add_table(
        doc,
        ["Стан", "Дозволена подія", "Новий стан", "Заборонені приклади"],
        [
            ("Draft", "видати або скасувати", "Active або Cancelled", "повернути, втратити"),
            ("Active", "прострочити, повернути, втратити", "Overdue, Returned, Lost", "видати повторно"),
            ("Overdue", "повернути або втратити", "Returned або Lost", "скасувати"),
            ("Returned", "немає", "кінцевий", "усі події"),
            ("Cancelled", "немає", "кінцевий", "усі події"),
            ("Lost", "немає", "кінцевий", "усі події"),
        ],
        [1.0, 2.35, 1.8, 1.65],
        8.0,
    )


def add_architecture(doc, diagram_path):
    new_page(doc)
    add_heading(doc, "3 Архітектура рішення")
    add_body(
        doc,
        "Рішення використовує три основні проєкти. Усередині LibraryDesk.Core домен, прикладний шар і інфраструктура відділені теками, просторами імен та портами. Домен не читає JSON, не звертається до Console і не створює конкретне сховище.",
    )
    add_image(doc, diagram_path, "Рисунок 1  Схема залежностей LibraryDesk", 6.65)
    add_table(
        doc,
        ["Шар", "Містить", "Заборонено"],
        [
            ("Domain", "Book, Reader, Loan, стани й правила", "Console, JSON, файли, сховища"),
            ("Application", "LibraryService, запити, результати, порти", "формат JSON і читання консолі"),
            ("Infrastructure", "JSON, системний час, файловий журнал", "правила строку та пені"),
            ("App", "меню, парсинг вводу, композиційний корінь", "розрахунки та переходи в обхід сервісу"),
        ],
        [1.15, 2.75, 2.9],
        8.2,
    )
    add_body(
        doc,
        "Архітектурні тести доводять, що LibraryDesk.Core не посилається на LibraryDesk.App, а публічний API простору Domain не відкриває типів Infrastructure. Це закріплює напрям залежностей у вибраній трьохпроєктній схемі.",
    )

    new_page(doc)
    add_heading(doc, "3.1 Композиційний корінь")
    add_listing(
        doc,
        "Лістинг 2  Створення залежностей у Program.cs",
        """IClock clock = new SystemClock();
ILibraryRepository repository = new JsonLibraryRepository(dataPath);
FileScenarioLogger scenarioLogger = new(logPath, clock);
SampleDataSeeder.Seed(repository);
LibraryService service = new(
    repository,
    clock,
    new LateFeePolicy(),
    scenarioLogger);
ConsoleApplication application = new(service, Console.In, Console.Out);
application.Run();""",
        8.1,
    )
    add_body(
        doc,
        "Конкретні реалізації створюються в одному місці. LibraryService отримує залежності через конструктор, тому тести підставляють InMemoryLibraryRepository, FixedClock і SpyScenarioLogger без зміни продуктового коду.",
    )
    add_heading(doc, "3.2 Точки розширення", 2)
    add_table(
        doc,
        ["Зміна", "Механізм", "Що не змінюється"],
        [
            ("CSV або база даних", "нова реалізація ILibraryRepository", "домен і сценарії"),
            ("Нове правило пені", "нова реалізація ILateFeePolicy", "Loan, меню, сховище"),
            ("Інший формат звіту", "окремий форматер рядків", "відбір і розрахунок"),
            ("Керований час", "нова реалізація IClock", "сервіс і сутності"),
        ],
        [2.0, 2.65, 2.15],
        8.6,
    )


def add_implementation(doc):
    new_page(doc)
    add_heading(doc, "4 Реалізація за практиками курсу")
    add_table(
        doc,
        ["Практика", "Місце в коді", "Підтвердження"],
        [
            ("decimal і дати", "Book, LoanItem, Loan", "гроші decimal; DateOnly і DateTimeOffset"),
            ("іменовані сталі", "LoanTermsPolicy, LateFeePolicy", "14 днів, день 8, множник 1,5"),
            ("короткі методи", "LibraryService, ConsoleApplication", "валідацію й форматування виділено"),
            ("DI", "LibraryService constructor", "4 реально замінні залежності"),
            ("барикада", "ValidateBook, ValidateReader, ValidateIssueRequest", "зовнішні дані перевіряються один раз"),
            ("власний виняток", "DomainRuleException", "заборонений перехід стану"),
            ("логування", "FileScenarioLogger", "Information, Warning, Error"),
            ("зберігання", "JsonLibraryRepository", "тимчасовий файл і атомарна заміна"),
        ],
        [1.45, 2.65, 2.7],
        7.8,
    )
    add_heading(doc, "4.1 Обробка помилок і журнал", 2)
    add_body(
        doc,
        "Очікувана помилка вводу повертається як Result<T>. Порушення життєвого циклу створеної сутності виражає DomainRuleException. ArgumentException та InvalidOperationException позначають помилку програміста. IOException журналюється на рівні Error і не проковтується.",
    )
    add_table(
        doc,
        ["Рівень", "Подія", "Приклад"],
        [
            ("Information", "початок і завершення UC", "Формуляр 1001 активовано"),
            ("Warning", "очікуване відхилення", "Читача або книгу не знайдено"),
            ("Error", "збій сховища або пошкоджені зв'язки", "IOException без вмісту файла"),
        ],
        [1.25, 2.35, 3.2],
        8.5,
    )
    add_heading(doc, "4.2 Файлове збереження", 2)
    add_body(
        doc,
        "JsonLibraryRepository відновлює доменні сутності через окремі серіалізаційні моделі, тому домен не містить JSON-атрибутів. Під час Save спочатку записується library.json.tmp, після чого файл замінюється. Інтеграційний тест зберігає книгу, читача й активний формуляр, відкриває новий екземпляр сховища та порівнює стан.",
    )

    new_page(doc)
    add_heading(doc, "4.3 Робота консольного застосунку")
    add_body(
        doc,
        "Замість демонстраційного сценарію реалізовано меню з шістьма пунктами. Словник команд замінює великий switch і тримає складність Run нижче межі. Некоректне число або дата не завершує процес: користувач отримує повідомлення і знову бачить меню.",
    )
    run_text = """LibraryDesk v1.0.0
1 — зареєструвати книгу
2 — зареєструвати читача
3 — видати книжку
4 — прийняти повернення
5 — знайти формуляри
6 — звіт «Боржники на дату»
0 — завершити"""
    add_listing(doc, "Лістинг 3  Перший екран застосунку", run_text, 9.0)
    add_body(
        doc,
        "Перший запуск генерує 50 книг і 20 умовних читачів. Робочі дані містяться у data/library.json, журнал — у data/librarydesk.log. Обидва файли виключено з Git, тому репозиторій не накопичує локальний стан або персональні дані.",
    )
    add_table(
        doc,
        ["Сценарій", "Ввід", "Вивід"],
        [
            ("Видача", "1001; 1; ISBN-0001", "строк повернення +14 днів"),
            ("Повернення", "1001; 2026-11-01", "8 днів; 72,00 грн"),
            ("Пошук", "читач 1; Active", "список за строком і ID"),
            ("Звіт", "2026-11-01", "боржники за спаданням днів"),
        ],
        [1.45, 2.25, 3.1],
        8.7,
    )


def add_testing(doc):
    new_page(doc)
    add_heading(doc, "5 Модульні тести та покриття")
    add_body(
        doc,
        "Основне рішення має 85 тестів, ізольований Legacy-модуль — 12. Набір перевищує мінімум 25 і включає параметризовану таблицю пені, переходи станів, валідацію, усі сценарії прикладного шару, пошук, звіт, JSON і архітектуру.",
    )
    add_table(
        doc,
        ["Група", "Мінімум", "Фактичне покриття прикладами"],
        [
            ("Пеня", "8", "10 випадків, межі 0, 7, 8, максимум і від'ємне"),
            ("Переходи", "5", "Draft, Active, Overdue, Returned, Cancelled, Lost"),
            ("Валідація", "4", "порожні рядки, 0, від'ємне, неправильний email"),
            ("Сценарії", "5", "UC-01–UC-06"),
            ("Пошук і звіт", "3", "порожній, один і кілька впорядкованих рядків"),
            ("Дублери", "2", "InMemoryLibraryRepository і FixedClock"),
        ],
        [1.55, 1.05, 4.2],
        8.4,
    )
    add_listing(
        doc,
        "Лістинг 4  Параметризовані випадки профільного правила",
        """[InlineData(0, 10, 1000, 0)]
[InlineData(1, 10, 1000, 10)]
[InlineData(7, 10, 1000, 70)]
[InlineData(8, 10, 1000, 120)]
[InlineData(10, 10, 1000, 150)]
[InlineData(8, 12.5, 1000, 150)]
[InlineData(20, 10, 300, 300)]""",
        8.6,
    )
    add_body(
        doc,
        "FixedClock фіксує 10.10.2026, тому тест UC-03 завжди очікує DueOn 24.10.2026. Фейкове сховище не торкається диска. SpyScenarioLogger дозволяє довести рівні Warning і Error без читання реального журналу.",
    )

    new_page(doc)
    add_heading(doc, "5.1 Результат покриття")
    add_image(doc, EVIDENCE / "coverage-summary.png", "Знімок екрана 1  Підсумкова сторінка покриття", 6.65)
    add_table(
        doc,
        ["Метрика", "Покрито", "Усього", "Відсоток", "Поріг"],
        [
            ("Рядки", "395", "463", "85,31 %", "70 %"),
            ("Гілки", "106", "148", "71,62 %", "інформаційно"),
        ],
        [1.4, 1.25, 1.2, 1.45, 1.5],
        8.8,
    )
    add_body(
        doc,
        "До знаменника входять Domain, Pricing, Services та Errors збірки LibraryDesk.Core. ConsoleApplication, JSON-адаптер, файловий журнал, obj і Legacy виключені явно у coverlet.runsettings. Це число характеризує правила та сценарії, а не кількість технічного коду.",
    )
    add_heading(doc, "5.2 Мутаційний експеримент", 2)
    add_body(
        doc,
        "Множник IncreasedRateMultiplier тимчасово змінено з 1,5 на 1,6. Із 10 профільних випадків почервоніли 3: 8 × 10,00; 10 × 10,00; 8 × 12,50. Фактичні значення 128,00, 160,00 і 160,00 не збіглися з очікуваними 120,00, 150,00 і 150,00. Після повернення 1,5 повна перевірка знову дала 97/97.",
    )

    new_page(doc)
    add_heading(doc, "5.3 Аналіз непокритих гілок", 2)
    add_table(
        doc,
        ["Місце", "Гілка", "Причина", "Рішення"],
        [
            ("LibraryService:100", "недодатний ID у UC-03", "не всі комбінації чисел", "TD-SR3-01, 0,5 год"),
            ("LibraryService:128", "недодатний ID у UC-04", "межова захисна перевірка", "TD-SR3-01, 0,5 год"),
            ("LibraryService:156", "недодатний фільтр", "парсер не передає таке значення", "лишити захист"),
            ("LibraryService:260", "формуляр не знайдено", "суміжне правило перевірено раніше", "TD-SR3-02, 0,5 год"),
            ("LibraryService:298", "ISBN відсутній у каталозі", "можливо лише після пошкодження JSON", "TD-SR3-03, 1 год"),
            ("Loan:81", "MarkOverdue до строку", "окрема команда не входить до меню 1.0", "лишити доменний захист"),
            ("ArchivedLoanSnapshot", "архівування", "історичний API ЛР-4", "TD-SR3-04, 2 год"),
        ],
        [1.35, 1.8, 2.25, 1.4],
        7.1,
    )
    add_body(
        doc,
        "У таблиці немає причини «не встиг». Непокриті місця класифіковано як недосяжні через штатну межу, свідомо поза сценаріями 1.0 або борг із конкретною оцінкою.",
    )


def add_ci(doc):
    new_page(doc)
    add_heading(doc, "6 Конвеєр безперервної інтеграції")
    add_body(
        doc,
        "Оновлений workflow запускається на кожен push і pull request до main. Він відновлює залежності, перевіряє dotnet format, збирає Release з TreatWarningsAsErrors, виконує обидва тестові проєкти, збирає Cobertura, перевіряє 70 % і публікує артефакт.",
    )
    ci_text = (ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
    add_listing(doc, "Лістинг 5  Актуальний workflow CI", ci_text, 6.6)
    add_body(
        doc,
        "Локальний скрипт tools/verify.sh виконує ті самі quality gates однією командою. Фактичний локальний запуск завершився успішно з 0 попереджень, 97 зеленими тестами та 85,31 % рядків.",
    )

    new_page(doc)
    add_heading(doc, "6.1 Доказ роботи критеріїв", 2)
    add_body(
        doc,
        "У ЛР-8 workflow вже було перевірено навмисним CS0168: червоний запуск завершився за 40 секунд, після виправлення зелений — за 45 секунд. Ці знімки підтверджують дію format/build/test gates попередньої версії. Для СР-3 поріг піднято з 60 до 70 % і додано окремий крок перевірки coverage.",
    )
    add_image(doc, LAB8_EVIDENCE / "ci-red.png", "Знімок екрана 2  Історичний червоний прогон quality gate", 6.55)
    add_image(doc, LAB8_EVIDENCE / "ci-green.png", "Знімок екрана 3  Історичний зелений прогон після виправлення", 6.55)
    add_body(
        doc,
        "На момент формування звіту sr03-individual-project існує лише локально. Після push треба зберегти зелений прогон саме для коміта версії 1.0.0, відкрити pull request і після злиття поставити тег v1.0. У документі цей зовнішній стан не позначено виконаним наперед.",
    )


def add_documentation(doc):
    new_page(doc)
    add_heading(doc, "7 Документація проєкту")
    add_body(
        doc,
        "README має вісім змістових розділів: робота, швидкий старт, структура, можливості, тести й покриття, технології, дані й резервна копія, обмеження. Новий користувач отримує застосунок трьома командами.",
    )
    add_listing(
        doc,
        "Лістинг 6  Три команди швидкого старту",
        """git clone https://github.com/kuum-oss/LibraryDesk.git && cd LibraryDesk
dotnet build LibraryDesk.sln -c Release
dotnet run --project LibraryDesk.App/LibraryDesk.App.csproj -c Release""",
        8.4,
    )
    add_table(
        doc,
        ["Артефакт", "Зміст"],
        [
            ("docs/spec.md", "межі, словник, 6 UC, таблиця пені, переходи"),
            ("docs/architecture.md", "шари, схема, composition root, точки розширення"),
            ("docs/user-guide.md", "запуск, 5 прохідних сценаріїв, помилки, backup"),
            ("docs/decisions.md", "7 рішень із альтернативами та мінусами"),
            ("docs/self-check.md", "20 пунктів, числа, слабкі місця й висновки"),
            ("docs/sr03/coverage-analysis.md", "межа, 7 гілок і мутаційний експеримент"),
        ],
        [2.4, 4.4],
        8.5,
    )
    add_heading(doc, "7.1 XML-документація", 2)
    add_body(
        doc,
        "GenerateDocumentationFile і TreatWarningsAsErrors застосовуються до всіх проєктів. Публічні типи, параметри, результати та навмисні винятки задокументовано; відсутній summary зупиняє Release-складання.",
    )


def add_decisions_and_self_check(doc):
    new_page(doc)
    add_heading(doc, "8 Журнал рішень і самооцінка")
    add_table(
        doc,
        ["№", "Рішення", "Головний плюс", "Прийнятий мінус"],
        [
            ("1", "JSON замість CSV", "природні вкладені позиції", "повний перезапис файла"),
            ("2", "Result плюс DomainRuleException", "різні класи помилок", "два способи обробки"),
            ("3", "три проєкти", "простіша збірка", "внутрішню межу контролює тест"),
            ("4", "агрегований repository", "узгоджений знімок", "ширший інтерфейс"),
            ("5", "IClock", "детерміновані тести", "додаткова залежність"),
            ("6", "без БД і GUI", "завершені сценарії", "однокористувацький режим"),
            ("7", "окрема межа coverage", "чесний знаменник", "інфраструктура рахується окремо"),
        ],
        [0.45, 2.15, 2.15, 2.05],
        7.6,
    )
    add_heading(doc, "8.1 Три підсумкові числа", 2)
    add_table(
        doc,
        ["Показник", "Значення"],
        [
            ("Пункти повністю", "17 із 20 — 85 %"),
            ("Покриття", "85,31 % рядків; 71,62 % гілок"),
            ("Залишковий борг", "7 годин без календарної вимоги"),
        ],
        [3.4, 3.4],
        9.1,
    )
    add_heading(doc, "8.2 Частково виконані пункти", 2)
    for text in [
        "Цикломатична складність не перевищує 8 за ручною перевіркою, але її межа ще не автоматизована в CI.",
        "Продуктове ядро не містить великих клонів, однак Legacy-модуль свідомо зберігає дефекти для попереднього аудиту.",
        "Комітів більше 25, але на 10.10.2026 історія охоплює 4 різні дні замість 6; дати не підроблялися.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "8.3 Три найслабші місця", 2)
    add_table(
        doc,
        ["Місце", "Наслідок", "Оцінка"],
        [
            ("Legacy і compatibility API", "можна випадково використати застарілий Order", "4 год"),
            ("Немає gate складності", "новий складний метод залишиться зеленим", "2 год"),
            ("Однофайловий JSON", "два процеси можуть перетерти зміни", "6 год для заміни"),
        ],
        [2.45, 3.15, 1.2],
        8.3,
    )


def add_conclusions(doc):
    new_page(doc)
    add_heading(doc, "9 Висновки")
    conclusions = [
        "Сформовано специфікацію з шістьма сценаріями, таблицею пені та життєвим циклом формуляра.",
        "Залежності спрямовано до домену, а конкретні JSON, час і журнал зібрано в композиційному корені.",
        "Профільне правило перевірено десятьма випадками; мутація множника зламала три тести.",
        "Локальний quality gate відтворює формат, Release, 97 тестів і поріг 70 % однією командою.",
        "Найважчим було визначити чесну межу покриття; підсумок 85,31 % стосується саме логіки.",
        "Найкориснішою практикою стала ін'єкція залежностей, бо фейкове сховище й FixedClock прибрали нестабільність.",
        "За повторного старту Domain, Application та Infrastructure доцільно одразу рознести по окремих збірках, а Legacy — по окремому рішенню.",
    ]
    for index, text in enumerate(conclusions, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        set_font(paragraph.add_run(f"{index}. {text}"), "Times New Roman", 10.5)

    add_heading(doc, "10 Використані джерела")
    sources = [
        "МакКоннелл С. Досконалий код. 2-ге видання. 2004.",
        "Мартін Р. Чистий код. 2008.",
        "Мартін Р. Чиста архітектура. 2017.",
        "Фаулер М. Рефакторинг. 2-ге видання. 2018.",
        "Microsoft Learn. Модульне тестування в .NET і XML-документація C#.",
        "GitHub Docs. Building and testing .NET та workflow artifacts.",
        "Документація xUnit, coverlet і ReportGenerator.",
    ]
    for source in sources:
        add_bullet(doc, source)
    add_heading(doc, "Поточний статус подання")
    add_body(
        doc,
        "Локальні артефакти СР-3 готові. Для завершення зовнішньої частини потрібно надіслати sr03-individual-project на GitHub, дочекатися зеленого CI, створити pull request до main, злити його та поставити тег v1.0. Після цього звіт слід оновити номером PR і знімком зеленого прогону саме версії 1.0.",
    )


def build_report():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    architecture = create_architecture_diagram()
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_contents(doc)
    add_intro(doc)
    add_specification(doc)
    add_architecture(doc, architecture)
    add_implementation(doc)
    add_testing(doc)
    add_ci(doc)
    add_documentation(doc)
    add_decisions_and_self_check(doc)
    add_conclusions(doc)
    doc.core_properties.title = "Самостійна робота 3 LibraryDesk"
    doc.core_properties.subject = "Індивідуальний проєкт конструювання застосунку за варіантом 2"
    doc.core_properties.author = "Гордєєв Дмитро Леонідович"
    doc.core_properties.keywords = "LibraryDesk, СР-3, .NET 8, тести, CI"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
