"""Generate the Ukrainian Lab 6 DOCX report from verified repository evidence."""

from pathlib import Path
import subprocess
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT / "tmp" / "docx" / "lab06"
OUT = ROOT / "output" / "docx" / "Звіт_ЛР6_LibraryDesk.docx"
FONT_DIR = Path("/System/Library/Fonts/Supplemental")

NAVY = "173F6D"
PALE_BLUE = "EEF4F8"
PALE_GRAY = "F5F6F7"
LIGHT_BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(85, 85, 85)


def run(command: list[str]) -> str:
    completed = subprocess.run(
        command,
        cwd=ROOT,
        check=True,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    return completed.stdout.strip()


def set_font(run_obj, name: str, size: float, bold: bool = False) -> None:
    run_obj.font.name = name
    run_obj.font.size = Pt(size)
    run_obj.font.bold = bold
    run_obj.font.color.rgb = BLACK
    run_obj._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run_obj._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run_obj._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = props.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        props.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = LIGHT_BORDER, width: str = "6") -> None:
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), width)
        node.set(qn("w:color"), color)


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def set_repeat_table_header(table) -> None:
    repeat_header(table.rows[0])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_obj = paragraph.add_run("Сторінка ")
    set_font(run_obj, "Arial", 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_obj._r.extend([begin, instruction, separate, placeholder, end])


def terminal_image(name: str, command: str, output: str, max_lines: int = 23) -> Path:
    lines = output.splitlines()[-max_lines:]
    wrapped: list[str] = []
    for line in lines:
        wrapped.extend(
            textwrap.wrap(
                line,
                width=92,
                subsequent_indent="  ",
                break_long_words=True,
                break_on_hyphens=False,
            )
            or [""]
        )
    width = 1500
    row_height = 31
    height = 94 + row_height * len(wrapped) + 26
    image = Image.new("RGB", (width, height), "#13202B")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((0, 0, width - 1, height - 1), radius=18, outline="#4D6577", width=2)
    draw.rounded_rectangle((0, 0, width - 1, 56), radius=18, fill="#243746")
    for index, color in enumerate(("#F27A72", "#F2CD72", "#73D19B")):
        draw.ellipse((25 + index * 29, 21, 40 + index * 29, 36), fill=color)
    font = ImageFont.truetype(str(FONT_DIR / "Andale Mono.ttf"), 21)
    draw.text((32, 70), "$ " + command, font=font, fill="#F4D88E")
    for index, line in enumerate(wrapped):
        draw.text((32, 70 + row_height * (index + 1)), line, font=font, fill="#E5EEF4")
    TMP.mkdir(parents=True, exist_ok=True)
    path = TMP / f"{name}.png"
    image.save(path)
    return path


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(document: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph(style="Body Text")
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_font(lead, "Times New Roman", 11, True)
    run_obj = paragraph.add_run(text)
    set_font(run_obj, "Times New Roman", 11)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table)
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_borders(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_obj = paragraph.add_run(header)
        set_font(run_obj, "Arial", 8, True)
        run_obj.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[column_index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if len(value) <= 18 or column_index == 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run_obj = paragraph.add_run(value)
            set_font(run_obj, "Arial", 7.7)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(document: Document, text: str, caption: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE_GRAY)
    set_cell_borders(cell)
    set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run_obj = paragraph.add_run(text.strip())
    set_font(run_obj, "Courier New", 7.8)
    add_caption(document, caption)


def add_screenshot(document: Document, image_path: Path, caption: str, width: float = 7.0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    add_caption(document, caption)


def page_break(document: Document) -> None:
    document.add_page_break()


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(21)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(18)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    body = document.styles["Body Text"]
    body.font.name = "Times New Roman"
    body.font.size = Pt(11)
    body.font.color.rgb = BLACK
    body.paragraph_format.first_line_indent = Inches(0.3)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    body.paragraph_format.space_after = Pt(6)

    caption = document.styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8)
    caption.font.italic = False
    caption.font.color.rgb = GRAY
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.text = "LibraryDesk   Лабораторна робота 6   "
    add_page_number(paragraph)


def build_document() -> None:
    build_output = run(["dotnet", "build", "LibraryDesk.sln", "--nologo"])
    test_output = run(
        [
            "dotnet",
            "test",
            "tests/Legacy.Tests/LibraryDesk.Legacy.Tests.csproj",
            "--no-build",
            "--nologo",
        ]
    )
    history_output = run(
        [
            "git",
            "log",
            "--reverse",
            "--oneline",
            "--grep=^refactor:",
            "main..HEAD",
        ]
    )
    build_image = terminal_image("build", "dotnet build LibraryDesk.sln", build_output, 18)
    test_image = terminal_image(
        "test",
        "dotnet test tests/Legacy.Tests/LibraryDesk.Legacy.Tests.csproj --no-build",
        test_output,
        16,
    )
    history_image = terminal_image(
        "history",
        "git log --reverse --oneline --grep=^refactor: main..HEAD",
        history_output,
        18,
    )

    document = Document()
    configure(document)
    document.core_properties.title = "Звіт до лабораторної роботи 6"
    document.core_properties.subject = "Рефакторинг успадкованого коду"
    document.core_properties.author = "Гордєєв Дмитро Леонідович"
    document.core_properties.keywords = "LibraryDesk, refactoring, characterization tests, .NET 8"

    # Page 1: title
    document.add_paragraph().paragraph_format.space_after = Pt(48)
    title = document.add_paragraph("ЗВІТ", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("до лабораторної роботи 6")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(subtitle.runs[0], "Arial", 13)
    topic = document.add_paragraph("Рефакторинг успадкованого коду за каталогом рефакторингів")
    topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic.paragraph_format.space_after = Pt(28)
    set_font(topic.runs[0], "Times New Roman", 14)
    for item in (
        "Дисципліна: Конструювання програмного забезпечення",
        "Проєкт: LibraryDesk",
        "Варіант: 2",
        "Виконав: Гордєєв Дмитро Леонідович",
        "Група: ІПЗ",
        "Гілка: lab6-refactoring",
    ):
        paragraph = document.add_paragraph(item)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.runs[0], "Times New Roman", 12)
    document.add_paragraph().paragraph_format.space_after = Pt(42)
    repo = document.add_paragraph("https://github.com/kuum-oss/LibraryDesk")
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(repo.runs[0], "Arial", 9)
    date = document.add_paragraph("2026")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(date.runs[0], "Times New Roman", 11)
    page_break(document)

    # Page 2: purpose and source data
    add_heading(document, "1 Мета роботи")
    add_body(
        document,
        "Мета роботи полягає у безпечній зміні внутрішньої будови успадкованого коду "
        "без зміни його зовнішньої поведінки. Незмінність доводиться 12 "
        "характеризаційними випадками, створеними до першого рефакторингу.",
    )
    add_heading(document, "2 Вихідні дані")
    add_body(
        document,
        "Універсальний модуль адаптовано до варіанта 2 LibraryDesk. Його додано "
        "окремим проєктом src/Legacy, щоб метрики не змішувалися з основним ядром.",
    )
    add_table(
        document,
        ["Універсальне ім'я", "Ім'я LibraryDesk", "Призначення"],
        [
            ["LegacyProcessor", "LegacyLoanProcessor", "Обробка документа видачі"],
            ["LineItem", "LoanLine", "Рядок документа"],
            ["Customer", "Reader", "Читач бібліотеки"],
            ["OrderRequest", "LoanRequest", "Параметри обробки"],
        ],
        [1.65, 2.05, 3.25],
    )
    add_body(
        document,
        "Початковий стан збережено комітом 1e2269a. Він навмисно містить "
        "дублювання, вкладені умови, тимчасові поля, магічні числа та рядкові стани.",
    )
    add_screenshot(document, build_image, "Знімок екрана 1  Збірка рішення без помилок і попереджень", 6.8)
    page_break(document)

    # Page 3: smells
    add_heading(document, "3 Реєстр дефектів проєктування")
    add_body(
        document,
        "Рядки прив'язані до початкового файлу в коміті 1e2269a. Дефекти "
        "впорядковано за впливом на безпечність наступних змін.",
    )
    add_table(
        document,
        ["№", "Дефект", "Місце", "Числовий прояв", "Рефакторинг"],
        [
            ["1", "Довгий метод", "Handle 27-126", "94 LOC, CC 17", "Extract Method"],
            ["2", "Вкладені умови", "Handle 40-68", "4 рівні", "Guard Clauses"],
            ["3", "Довгий список", "Handle 27-38", "11 параметрів", "Parameter Object"],
            ["4", "Дублювання", "Handle, Preview", "14 рядків", "Спільні методи"],
            ["5", "Великий клас", "LegacyLoanProcessor", "7 відповідальностей", "Extract Class"],
            ["6", "Primitive obsession", "вид і стан", "7 порівнянь", "Enums"],
            ["7", "Feature envy", "DescribeReader", "5 полів Reader", "Move Method"],
            ["8", "Temporary field", "рядки 23-24", "2 поля", "Remove Field"],
            ["9", "Магічні числа", "формули", "7 значень", "Constants"],
            ["10", "Мертві параметри", "Handle", "2 параметри", "Parameter Object"],
        ],
        [0.3, 1.35, 1.45, 1.45, 2.35],
    )
    add_body(
        document,
        "Найвищий ризик створював Handle: він одночасно перевіряв дані, рахував "
        "вартість і доставку, будував звіт та змінював журнал.",
    )
    page_break(document)

    # Page 4: tests
    add_heading(document, "4 Характеризаційні тести")
    add_body(
        document,
        "До рефакторингів створено 12 випадків: чотири коди помилок, золотий "
        "зразок звіту, чотири варіанти знижки, межа доставки, опис читача і журнал.",
    )
    add_code(
        document,
        '''[Theory]
[InlineData(null, LoanState.New, "a@b.c", "ERR: null")]
[InlineData("empty", LoanState.New, "a@b.c", "ERR: empty")]
[InlineData("two", LoanState.Draft, "a@b.c", "ERR: state")]
[InlineData("two", LoanState.New, "no-mail", "ERR: mail")]
public void Handle_BadInput_ReturnsErrorCode(...)
{
    var actual = sut.Handle(request);
    Assert.Equal(expected, actual);
}''',
        "Лістинг 1  Фіксація чотирьох помилкових гілок",
    )
    add_code(
        document,
        '''var expected =
    "Документ #1001\\n" +
    "Клієнт: Іван\\n" +
    "A-1 x2 = 300.00 UAH\\n" +
    "B-2 x1 = 400.00 UAH\\n" +
    "Знижка: 105.00 UAH\\n" +
    "Доставка: 60.00 UAH\\n" +
    "Разом: 655.00 UAH\\n";
Assert.Equal(expected, report);''',
        "Лістинг 2  Золотий зразок повного звіту",
    )
    add_screenshot(document, test_image, "Знімок екрана 2  Дванадцять тестових випадків до і після рефакторингу", 6.8)
    page_break(document)

    # Page 5: metrics before
    add_heading(document, "5 Метрики до рефакторингу")
    add_body(
        document,
        "LOC пораховано як непорожні фізичні рядки без рядків, що містять лише "
        "коментар. CC дорівнює одиниці плюс кількість точок розгалуження.",
    )
    add_table(
        document,
        ["Метод", "LOC", "CC", "Вкладеність", "Параметрів"],
        [
            ["Handle", "94", "17", "4", "11"],
            ["Preview", "39", "9", "2", "3"],
            ["DescribeReader", "16", "3", "1", "1"],
            ["DumpLog", "9", "2", "1", "0"],
            ["Клас загалом", "164", "-", "-", "-"],
        ],
        [2.1, 0.8, 0.8, 1.45, 1.45],
    )
    add_heading(document, "5 1 Розрахунок CC методу Handle", 2)
    add_table(
        document,
        ["Точка розгалуження", "Кількість"],
        [
            ["items та Count", "2"],
            ["стан і оператор ||", "2"],
            ["пошта і оператор &&", "2"],
            ["два цикли", "2"],
            ["три види читача", "3"],
            ["три межі знижки", "3"],
            ["межа доставки", "1"],
            ["sendMail", "1"],
            ["Разом точок", "16"],
        ],
        [5.8, 1.15],
    )
    add_body(document, "CC = 1 + 16 = 17. Це значно вище за орієнтир курсу CC <= 5.")
    page_break(document)

    # Page 6: refactoring log
    add_heading(document, "6 Журнал рефакторингів")
    add_body(
        document,
        "Кожен крок має власний refactor-коміт. Після кожного коміту виконано "
        "ті самі 12 випадків; усі запуски завершилися успішно.",
    )
    add_table(
        document,
        ["№", "Хеш", "Рефакторинг", "Дефект", "Тести"],
        [
            ["1", "09fab10", "Magic Number to Constant", "магічні числа", "12/12"],
            ["2", "42e47ab", "Primitive to Enum", "рядкові стани", "12/12"],
            ["3", "155e552", "Guard Clauses", "вкладені умови", "12/12"],
            ["4", "d8bc07c", "Extract Method", "довгий метод", "12/12"],
            ["5", "2360424", "Shared Methods", "дублювання", "12/12"],
            ["6", "74ea106", "Extract Class", "політики", "12/12"],
            ["7", "ab10d9b", "ReportBuilder", "форматування", "12/12"],
            ["8", "a2cb51b", "Parameter Object", "11 параметрів", "12/12"],
            ["9", "928c47e", "Move Method", "feature envy", "12/12"],
            ["10", "fa57ab7", "Remove Field", "2 temp fields", "12/12"],
            ["11", "bc4c327", "Polymorphism", "умовний вибір", "12/12"],
        ],
        [0.3, 0.85, 2.3, 2.0, 0.7],
    )
    add_screenshot(document, history_image, "Знімок екрана 3  Одинадцять окремих refactor комітів", 6.8)
    page_break(document)

    # Page 7: before and after
    add_heading(document, "7 Порівняння було і стало")
    add_heading(document, "7 1 Перевірка вхідних даних", 2)
    add_code(
        document,
        '''// БУЛО: чотири рівні вкладеності
if (items != null)
  if (items.Count > 0)
    if (state == "new" || state == "paid")
      if (readerMail != null && readerMail.Contains("@"))
        _log.Add("ok " + docId);
      else return "ERR: mail";
    else return "ERR: state";
  else return "ERR: empty";
else return "ERR: null";''',
        "Лістинг 3  Вкладена перевірка до рефакторингу",
    )
    add_code(
        document,
        '''// СТАЛО: порядок помилок збережено
if (request.Items is null) return "ERR: null";
if (request.Items.Count == 0) return "ERR: empty";
if (request.State is not (LoanState.New or LoanState.Paid))
    return "ERR: state";
if (request.Reader.Email is null ||
    !request.Reader.Email.Contains("@"))
    return "ERR: mail";
return null;''',
        "Лістинг 4  Охоронні умови після рефакторингу",
    )
    add_body(
        document,
        "Порядок перевірок не змінено, тому дані з кількома порушеннями повертають "
        "той самий перший код. Максимальна вкладеність зменшилася з чотирьох рівнів до одного.",
    )
    page_break(document)

    # Page 8: parameter object and extracted classes
    add_heading(document, "8 Об'єкт параметрів і виділені класи")
    add_code(
        document,
        '''// БУЛО: 11 параметрів
public string Handle(int docId, int readerId,
    string readerName, string? readerMail,
    ReaderKind readerKind, int readerDone,
    List<LoanLine>? items, LoanState state,
    string currency, DateTime createdAt,
    bool sendMail)''',
        "Лістинг 5  Сигнатура Handle до Introduce Parameter Object",
    )
    add_code(
        document,
        '''public sealed record LoanRequest(
    int DocumentId,
    Reader Reader,
    IReadOnlyList<LoanLine>? Items,
    LoanState State,
    string Currency,
    bool SendMail);

public string Handle(LoanRequest request)''',
        "Лістинг 6  Один параметр після рефакторингу",
    )
    add_body(
        document,
        "Параметри readerId і createdAt не використовувалися, тому їх не перенесено. "
        "Знижка, доставка й текст звіту винесені до DiscountPolicy, ShippingPolicy і "
        "ReportBuilder. Метод опису перенесено до Reader.",
    )
    add_table(
        document,
        ["Тип", "Відповідальність", "Причина зміни"],
        [
            ["LegacyLoanProcessor", "Координація і журнал", "Один сценарій"],
            ["DiscountPolicy", "Розмір знижки", "Правила змінюються разом"],
            ["ShippingPolicy", "Вартість доставки", "Окрема бізнес-політика"],
            ["ReportBuilder", "Формат тексту", "Золотий зразок локалізує ризик"],
            ["IDiscountRule", "Поліморфні ставки", "Нове правило без зміни політики"],
        ],
        [1.8, 2.15, 3.0],
    )
    page_break(document)

    # Page 9: metrics after
    add_heading(document, "9 Метрики після рефакторингу")
    add_table(
        document,
        ["Показник", "До", "Після", "Зміна"],
        [
            ["LOC Handle", "94", "19", "-79,8 %"],
            ["CC Handle", "17", "3", "-82,4 %"],
            ["Вкладеність Handle", "4", "1", "-3 рівні"],
            ["Параметри Handle", "11", "1", "-10"],
            ["Дубльовані рядки", "14", "0", "-14"],
            ["Магічні числа у формулах", "7", "0", "-7"],
            ["Тимчасові поля", "2", "0", "-2"],
            ["Відповідальності процесора", "7", "2", "-5"],
            ["LOC модуля", "181", "227", "+25,4 %"],
        ],
        [3.1, 0.9, 0.9, 1.4],
    )
    add_body(
        document,
        "Загальний LOC зріс через об'єкт параметрів, політики, будівник звіту й "
        "поліморфні правила. Це очікувано: головний метод скоротився майже на 80 %, "
        "а зміни формату, знижки та доставки тепер ізольовані.",
    )
    add_heading(document, "9 1 Доказ незмінності поведінки", 2)
    add_body(
        document,
        "До рефакторингу 12 випадків були зеленими. Після одинадцяти комітів ті "
        "самі 12 випадків із незміненими очікуваннями зелені. Кожен проміжний "
        "refactor-коміт перевірено в окремому worktree; червоних комітів немає. "
        "Отже, зовнішня поведінка модуля збережена.",
    )
    page_break(document)

    # Page 10: technical debt
    add_heading(document, "10 Реєстр технічного боргу")
    add_table(
        document,
        ["ID", "Опис", "Вид", "Год", "Грн", "Пріоритет"],
        [
            ["TD-01", "Формат чисел залежить від культури", "необачний код", "1", "300", "високий"],
            ["TD-02", "Опис читає системний годинник", "необачний код", "2", "600", "середній"],
            ["TD-03", "Коди помилок є рядками", "свідома архітектура", "3", "900", "середній"],
            ["TD-04", "Публічні змінювані поля", "необачний код", "2", "600", "високий"],
            ["TD-05", "Ставки зашиті у збірку", "свідомий код", "3", "900", "низький"],
            ["TD-06", "Тести змінюють CurrentCulture", "свідомі тести", "1,5", "450", "середній"],
            ["Разом", "", "", "12,5", "3750", ""],
        ],
        [0.65, 2.8, 1.4, 0.55, 0.6, 0.95],
    )
    add_heading(document, "10 1 Вартість відсотків", 2)
    add_body(
        document,
        "Модуль змінюють приблизно 4 години на місяць. TD-01, TD-02 і TD-04 "
        "збільшують тривалість змін на 25 %, тобто на 1 годину або 300 грн на "
        "місяць. Погашення TD-01 і TD-04 коштує 3 години, або 900 грн, і "
        "окупиться приблизно за 3 місяці.",
    )
    add_body(
        document,
        "Борг залишено навмисно: InvariantCulture і керований годинник змінять "
        "зафіксовану поведінку та потребують окремих вимог і тестів.",
    )
    page_break(document)

    # Page 11: conclusions
    add_heading(document, "11 Висновки")
    conclusions = [
        "Метод Handle скорочено з 94 до 19 LOC, а CC — з 17 до 3.",
        "Одинадцять параметрів замінено одним LoanRequest; два мертві параметри не перенесено.",
        "Найризикованішим був ReportBuilder, бо золотий зразок фіксує кожен пробіл і перенос рядка.",
        "Одинадцять малих комітів перевірено однаковими 12 характеризаційними випадками.",
        "Загальний LOC зріс на 25,4 %, але складність головного методу зменшилася на 82,4 %.",
        "Залежність від культури та годинника залишено в боргу, бо її виправлення змінює контракт.",
    ]
    for index, conclusion in enumerate(conclusions, 1):
        paragraph = document.add_paragraph(style="Body Text")
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        run_obj = paragraph.add_run(f"{index}. {conclusion}")
        set_font(run_obj, "Times New Roman", 11)
    add_heading(document, "11 1 Підсумок перевірки", 2)
    add_table(
        document,
        ["Критерій", "Результат"],
        [
            ["Реєстр дефектів", "10 типів з рядками й числами"],
            ["Характеризаційні тести", "12 випадків до рефакторингів"],
            ["Серія рефакторингів", "11 окремих refactor-комітів"],
            ["Проміжна історія", "11 зелених комітів"],
            ["Збірка", "0 warnings, 0 errors"],
            ["Технічний борг", "6 пунктів, 12,5 год, 3750 грн"],
        ],
        [3.2, 3.75],
    )
    add_body(document, "Репозиторій: https://github.com/kuum-oss/LibraryDesk")
    add_body(document, "Гілка: lab6-refactoring")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
